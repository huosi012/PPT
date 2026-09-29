# -*- coding: utf-8 -*-
"""deck_components —— 建立在 pptkit 之上的页级组件（工程场景条、要点卡、附图组合、占位框等）。

用法：build 脚本里 `from pptkit import *` 之后再 `from deck_components import *`。
这些组件在一次课题申报 PPT 中打磨过：投影上只放结论与要点，细节写进备注；附图整组可替换。
"""
from pptx.enum.shapes import MSO_SHAPE  # noqa: F401  (build 脚本画附图时常用)
from pptx.oxml.ns import qn

from pptkit import *  # noqa: F401,F403

BAR = STYLE['top_bar']  # 卡片顶线 / 左侧色条粗细（随配色方案变化：现版 6px，主红·留白 3px）

FIG_NOTE = ('说明：“现状与问题”可替换为本单位的实际案例；“示意图”为临时附图（形状组合，可直接编辑），'
            '如有实际图片，选中该组合删除后插入即可。')


# ---------------------------------------------------------------- 左栏 / 场景条 / 要点卡
def content_left(s, layer, label=None, body=None, outs=(), outs_bottom=None):
    """研究内容页左栏：四层导航 +（可选）补充说明 + 产出结论条。返回底部 y。
    outs_bottom 给定时产出块贴底，底边与右侧附图对齐（左栏只放导航和产出时用）。"""
    th = LAYER[layer]
    if HEADER_STYLE == 'banner':      # 页眉已写明是哪一层：导航只写层级名，四条两两对齐场景条（CT—CT+104）与要点卡（CT+116—CT+212）
        bars = [(CONTENT_TOP, 48), (CONTENT_TOP + 56, 48), (CONTENT_TOP + 116, 44), (CONTENT_TOP + 168, 44)]
        for (by, bh), i in zip(bars, (4, 3, 2, 1)):
            t = LAYER[i]
            on = (i == layer)
            rect(s, 64, by, 260, bh, fill=t.main if on else WHITE, line=None if on else LINE)
            if not on:
                rect(s, 64, by, STYLE['card_bar'], bh, fill=t.main)
            text(s, 84, by, 230, bh, LAYER_NO[i] + ' ' + LAYER_KIND[i], size=19, bold=True, color=WHITE if on else SUB,
                 anchor='m', lh=26, wrap=False)
        y = CONTENT_TOP + 212 + 22
    else:
        y = layer_nav(s, 64, 214, 260, layer) + 22
    if label:
        text(s, 64, y, 260, 34, label, size=22, bold=True, color=INK, lh=32, wrap=False)
        paras = body if isinstance(body, list) else [{'t': body}]
        bh = sum(wrap_lines(p['t'], 260, 19) * 30 + p.get('sa', 0) for p in paras)
        text(s, 64, y + 46, 260, bh + 4, body, size=19, color=SUB, lh=30)
        y += 46 + bh + 28
    if outs:
        hs = [wrap_lines('✓ ' + o, 260 - STYLE['concl_bar'] - 24, 18, True) * 26 + 22 for o in outs]
        if outs_bottom:
            y = max(y, outs_bottom - 44 - sum(hs) - 8 * (len(hs) - 1))
        text(s, 64, y, 260, 34, '产出', size=22, bold=True, color=INK, lh=32, wrap=False)
        y += 44
        for o, h in zip(outs, hs):
            concl(s, 64, y, 260, h, o, th, size=18, check=True, pad=12)
            y += h + 8
    return y


def scene_band(s, x, y, w, h, scene, role, scene_label='工程场景', role_label='本层作用', split=0.6):
    """场景条：左为实际工程场景（现状与问题），右为本页要做什么。放在研究内容页顶部。"""
    rect(s, x, y, w, h, fill=WHITE, line=LINE)
    rect(s, x, y, STYLE['concl_bar'], h, fill=RED)
    lw = round(w * split)
    tag(s, x + 24, y + 12, scene_label, fill=RED, size=16, h=28, padx=10)
    text(s, x + 24, y + 46, lw - 52, h - 50, scene, size=18, color=INK, lh=27)
    line(s, x + lw, y + 16, x + lw, y + h - 16, RULE, 1)
    text(s, x + lw + 28, y + 12, 240, 28, role_label, size=16, bold=True, color=MUTED, anchor='m', lh=22, wrap=False)
    text(s, x + lw + 28, y + 46, w - lw - 56, h - 50, role, size=18, color=BODY, lh=27)


def card_title(s, x, y, w, title, th, size=25, uw=None):
    """居中标题 + 同色下划线。"""
    text(s, x, y, w, 40, title, size=size, bold=True, color=INK, align='c', anchor='m', lh=round(size * 1.3),
         wrap=False)
    uw = uw or min(w - 40, text_width(title, size, True) + 12)
    rect(s, x + (w - uw) / 2.0, y + 48, uw, 3, fill=th.main)


def point_card(s, x, y, w, h, no, title, desc, th, sub=None, title_size=20, desc_size=17):
    """要点卡：顶部色条 + 编号标题（+ 副标题）+ 说明。一页放 3—5 张。"""
    rect(s, x, y, w, h, fill=WHITE, line=LINE)
    rect(s, x, y, w, BAR, fill=th.main)
    text(s, x + 22, y + 16, w - 44, 32, '<c=%s><m>%s</m></c>  %s' % (th.text, no, title), size=title_size, bold=True,
         color=INK, anchor='m', lh=round(title_size * 1.3), wrap=False)
    ty = y + 58
    if sub:
        text(s, x + 22, ty - 4, w - 44, 24, sub, size=15, bold=True, color=th.text, lh=22, wrap=False)
        ty += 26
    text(s, x + 22, ty, w - 44, y + h - ty - 8, desc, size=desc_size, color=BODY, lh=round(desc_size * 1.52))


def big_numbers(s, x, y, w, items, h=68, gap=12, label_w=70):
    """一行大数字：items = [(数值, 单位, 标签)]。用于“三个数字”“工程现状”。"""
    tw = (w - gap * (len(items) - 1)) / float(len(items))
    for i, (v, u, lab_) in enumerate(items):
        xx = x + i * (tw + gap)
        rect(s, xx, y, tw, h, fill=WHITE, line=LINE)
        text(s, xx + 16, y, tw - 32, h, '%s<s=16><n><c=%s> %s</c></n></s>' % (v, SUB, u), size=32, bold=True,
             color=INK, font=MONO, anchor='m', lh=40, wrap=False)
        text(s, xx + tw - 16 - label_w, y, label_w, h, lab_, size=16, bold=True, color=MUTED, align='r', anchor='m',
             lh=22, wrap=False)


def target_chip(s, x, y, label, value, size=17, h=34, padx=12):
    """量化目标占位：虚线框 + 指标名 + 目标值（如“≥【待填】%”）。返回宽度。"""
    w = text_width(label + '　', size) + text_width(value, size, True) + padx * 2
    shp = rect(s, x, y, w, h, fill=PH.bg, line=PH.main, lw=1.2, dash='dash')
    text(s, x, y, w, h, '%s　<bc=%s>%s</bc>' % (label, PH.text, value), size=size, color=BODY, align='c', anchor='m',
         lh=round(size * 1.3), wrap=False, shp=shp)
    return w


# ---------------------------------------------------------------- 附图（形状组合，整组可替换）
class Fig:
    """附图容器：所有形状收进一个组合；用户选中组合删除，插入真图即可。"""

    def __init__(self, s, name):
        self.grp = s.shapes.add_group_shape()
        self.grp.name = name
        self.shapes = self.grp.shapes
        self._no = s._no


def fig_area(s, x, y, w, h, title, note='示意图'):
    """附图区：小节标题 + 浅底面板（面板与图内形状同属一个组合）。返回 (组合, 面板顶 y)。"""
    header(s, x, y, title, note=note, w=w)
    f = Fig(s, '附图 · ' + strip_tags(title))
    rect(f, x, y + 46, w, h - 46, fill=PANEL)
    return f, y + 46


def fig_panel(s, x, y, w, h, title):
    """附图面板（不画小节标题）：返回组合，面板与图内形状同属一个组合。"""
    f = Fig(s, '附图 · ' + strip_tags(title))
    rect(f, x, y, w, h, fill=PANEL)
    return f


def fit_fig(f, y0, h0, y1, h1):
    """把按旧尺寸（面板顶 y0、高 h0）画好的附图组合整体搬到新位置并按高度压缩（宽度不变）。
    组合内图片（图标）先反向拉高，压缩后仍保持原比例。"""
    k = h1 / float(h0)
    g = f.grp._element
    xf = g.grpSpPr.find(qn('a:xfrm'))
    off, ext, choff, chext = (xf.find(qn(t)) for t in ('a:off', 'a:ext', 'a:chOff', 'a:chExt'))
    for pic in g.iter(qn('p:pic')):
        px = pic.find(qn('p:spPr')).find(qn('a:xfrm'))
        o, e = px.find(qn('a:off')), px.find(qn('a:ext'))
        h = int(e.get('cy'))
        h2 = int(round(h / k))
        o.set('y', str(int(o.get('y')) - (h2 - h) // 2))
        e.set('cy', str(h2))
    cy0, ch = int(choff.get('y')), int(chext.get('cy'))
    off.set('y', str(int(round(E(y1) + (cy0 - E(y0)) * k))))
    ext.set('cy', str(int(round(ch * k))))


def fig_placeholder(s, x, y, w, h, title, hint):
    """预留附图位置：虚线框 + 图片图标 + 名称 + 建议内容。"""
    f = Fig(s, '附图位置 · ' + title)
    rect(f, x, y, w, h, fill=PH.bg, line=PH.main, lw=1.2, dash='dash')
    icon(f, 'photo', PH.main, x + (w - 40) / 2.0, y + h / 2.0 - 54, 40)
    text(f, x + 16, y + h / 2.0 - 6, w - 32, 28, '附图位置 · ' + title, size=18, bold=True, color=PH.text, align='c',
         lh=26, wrap=False)
    text(f, x + 16, y + h / 2.0 + 24, w - 32, 24, hint, size=15, color=PH.text, align='c', lh=22, wrap=False)
    return f


def fnode(f, x, y, w, h, txt, th, kind='n', size=17):
    """附图节点：n 普通 / k 强调 / on 实心 / w 警示。"""
    st = {'n': dict(color=BODY, fill=WHITE, line=LINE),
          'k': dict(color=th.text, fill=th.bg, line=th.soft, bold=True),
          'on': dict(color=on_color(th.main), fill=th.main, line=None, bold=True),
          'w': dict(color=RED, fill=RED_BG, line=RED, bold=True)}[kind]
    return chip(f, x, y, txt, w=w, h=h, size=size, **st)


def arrow(f, x1, y1, x2, y2, color=MUTED, lw=2, dash=None):
    return line(f, x1, y1, x2, y2, color, lw, dash=dash, tail='triangle')


def smooth(pts, n=8):
    """Catmull-Rom 插值：把折线点加密成平滑曲线（画示意曲线用）。"""
    P = [pts[0]] + list(pts) + [pts[-1]]
    out = []
    for i in range(1, len(P) - 2):
        p0, p1, p2, p3 = P[i - 1], P[i], P[i + 1], P[i + 2]
        for k in range(n):
            t = k / float(n)
            out.append(tuple(0.5 * (2 * p1[j] + (p2[j] - p0[j]) * t + (2 * p0[j] - 5 * p1[j] + 4 * p2[j] - p3[j]) * t * t
                                    + (3 * p1[j] - p0[j] - 3 * p2[j] + p3[j]) * t ** 3) for j in (0, 1)))
    out.append(tuple(pts[-1]))
    return out
