# -*- coding: utf-8 -*-
"""pptkit —— 复刻参考页版式的 python-pptx 组件库。

坐标系统一使用 1920×1080 像素（16:9 宽屏，1px = 6350 EMU = 0.5pt），
字号参数同样以像素计（22px = 11pt）。
"""
import os
import re
import json
import subprocess

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ICON_DIR = os.path.join(ROOT, 'assets', 'icons')
TOOLS_DIR = os.path.join(ROOT, 'tools')

W, H = 1920, 1080
PX = 6350


def E(v):
    return Emu(int(round(v * PX)))


# ---------------------------------------------------------------- 字体
FONT = '微软雅黑'      # 中文/正文：Windows、Mac Office 均自带
MONO = 'Consolas'      # 编号、页码、数字
SYMBOL = 'Arial'       # 项目符号三角

# ---------------------------------------------------------------- 色板（取自参考页）
INK = '0F1A2B'        # 标题、粗体
BODY = '3A4556'       # 正文
SUB = '6B7785'        # 副标题、说明
MUTED = '7C889A'      # 标签说明、页脚（对比度约 3.6:1）
LINE = 'E3E7EF'       # 卡片描边
RULE = 'D3D9E1'       # 页眉页脚细线
GRAYBG = 'F3F5F7'     # 问题卡底
PANEL = 'F7F9FC'      # 浅面板底
GRAYBAR = '9AA6B8'    # 问题卡左侧灰条
TRI = 'C3CCD9'        # 流向三角
RULER = 'A8B4C6'      # 标尺线
WHITE = 'FFFFFF'
RED = 'C00000'
RED_BG = 'FDF2F2'


class Theme:
    def __init__(self, main, bg, soft, num, text):
        self.main, self.bg, self.soft, self.num, self.text = main, bg, soft, num, text


# ---------------------------------------------------------------- 配色方案
# 颜色按“用途”取用：M1–M5 模块序号色；LAYER 四层架构色；OK 正向（目标/回流/已解决）；PH 待填占位；
# STAGE 四个阶段；ARCH 架构分区标签；KPI 重点数字；BUDGET2 经费“人员”段；LOCO 机车插画。
# 通过环境变量 PPT_PALETTE 切换：vivid（初版，参考页原色）/ formal（稳重蓝）/ muted（低饱和多色）。
T = Theme
_VIVID = dict(
    M=[T('0A4CFF', 'F4F7FF', 'EAF0FF', 'D6E1FD', '0A4CFF'), T('7C5CF0', 'F5F2FE', 'ECE6FD', 'DDD4FB', '6A4ADE'),
       T('00A37A', 'EFF9F6', 'DDF3EC', 'C6EADF', '0B8A67'), T('D97706', 'FEF5EA', 'FCEBD3', 'F7DDBA', 'B45F04'),
       T('0098D4', 'EEF8FD', 'DDF1FB', 'C4E8F7', '0284C7')],
    LAYER=('M4', 'M1', 'M2', 'M3'),
    OK='M3', PH='M4',
    STAGE=['92C34D', 'FFC000', '00B0F0', '7030A0'],
    STAGE_ON=[WHITE, INK, WHITE, WHITE], STAGE_BULLET=['92C34D', 'D9A300', '00B0F0', '7030A0'],
    ARCH='7030A0', KPI=RED, BUDGET2='4A78FF', LOCO='locomotive',
)
_NAVY = T('1F4E79', 'F3F6FA', 'E3EAF3', 'D3DDE9', '1F4E79')
_FORMAL = dict(
    M=[_NAVY] * 5,
    LAYER=(T('5F7EA5', 'F4F7FA', 'E5ECF4', 'D7E0EC', '4F6C90'), T('3F6A9E', 'F2F6FA', 'E2EAF4', 'D2DDEB', '355D8E'),
           T('2A5486', 'F1F5F9', 'E0E8F1', 'D0DBE8', '2A5486'), T('1B3B63', 'F0F3F7', 'DDE4ED', 'CDD6E2', '1B3B63')),
    OK=T('2E7263', 'F1F7F5', 'DFEDE9', 'CFE3DD', '2A6658'),
    PH=T('A36A1E', 'FBF5EC', 'F5E8D2', 'EEDCBD', '8F5C18'),
    STAGE=['9DB2CC', '6D8CB3', '3F6A9E', '1F4371'],
    ARCH='1F4E79', KPI='1F4E79', BUDGET2='5B7FAE', LOCO='locomotive_formal',
)
_MUTED = dict(
    M=[T('2F5D8A', 'F2F6FA', 'E2EAF3', 'D2DDEA', '2F5D8A'), T('5B5E91', 'F5F5F9', 'E7E7F0', 'D9DAE8', '4F5285'),
       T('2F7564', 'F1F7F5', 'DEEDE8', 'CCE2DC', '2A6859'), T('9A6A2E', 'FBF6EE', 'F3E8D6', 'EBDBC2', '85591F'),
       T('2F7892', 'F0F6F9', 'DDEBF1', 'CCE0E9', '296A81')],
    LAYER=('M4', 'M1', 'M2', 'M3'),
    OK='M3', PH=T('A36A1E', 'FBF5EC', 'F5E8D2', 'EEDCBD', '8F5C18'),
    STAGE=['7E9A5E', 'C3A04A', '4F8AA6', '6A5A8C'],
    ARCH='5B5E91', KPI='2F5D8A', BUDGET2='6A8BB2', LOCO='locomotive_formal',
)
PALETTES = {'vivid': _VIVID, 'formal': _FORMAL, 'muted': _MUTED}
PALETTE = os.environ.get('PPT_PALETTE', 'vivid')
_P = PALETTES[PALETTE]


def _pick(v):
    return _P['M'][int(v[1]) - 1] if isinstance(v, str) else v


M1, M2, M3, M4, M5 = _P['M']
OK = _pick(_P['OK'])
PH = _pick(_P['PH'])
STAGE = list(_P['STAGE'])
STAGE_ON = _P.get('STAGE_ON')          # 阶段色块上的文字色（缺省按亮度自动取）
STAGE_BULLET = _P.get('STAGE_BULLET')  # 阶段卡片项目符号色（缺省自动加深）
ARCH = _P['ARCH']
KPI_COLOR = _P['KPI']
BUDGET2 = _P['BUDGET2']
LOCO_NAME = _P['LOCO']
# 兼容旧名称
BLUE, PURPLE, GREEN, ORANGE, CYAN = M1, M2, M3, M4, M5
REDT = Theme('C00000', 'FDF2F2', 'F9E1E1', 'F1CACA', 'C00000')
GRAY = Theme('9AA6B8', 'F3F5F7', 'E9EDF2', 'DCE2EA', '6B7785')

# 四层架构的固定配色：①基础层 ②决策层 ③执行层 ④集成层
LAYER = {i + 1: _pick(v) for i, v in enumerate(_P['LAYER'])}


def luminance(hexcol):
    def ch(c):
        c = c / 255.0
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = (int(hexcol[i:i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * ch(r) + 0.7152 * ch(g) + 0.0722 * ch(b)


def on_color(hexcol):
    """色块上的文字颜色：深色底用白字，浅色底用深墨字。"""
    return WHITE if 1.05 / (luminance(hexcol) + 0.05) >= 3.8 else INK


def ink_safe(hexcol):
    """白底上作为线条/符号时可辨的颜色：过浅则向深墨色混合加深。"""
    if luminance(hexcol) <= 0.3:
        return hexcol
    a = [int(hexcol[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(INK[i:i + 2], 16) for i in (0, 2, 4)]
    return '%02X%02X%02X' % tuple(round(x * 0.62 + y * 0.38) for x, y in zip(a, b))


LAYER_NO = {1: '①', 2: '②', 3: '③', 4: '④'}
LAYER_KIND = {1: '基础层', 2: '决策层', 3: '执行层', 4: '集成层'}
LAYER_NAME = {1: '测试知识库与任务工作流', 2: '任务理解与智能编排调度', 3: '智能体执行', 4: '架构融合与协同运行'}

# ---------------------------------------------------------------- 文本测量（用于溢出预警）
_FONT_FILES = {
    False: '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
    True: '/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc',
}
_pil_cache = {}


def _pil(bold):
    if bold not in _pil_cache:
        try:
            _pil_cache[bold] = ImageFont.truetype(_FONT_FILES[bold], 100, index=2)
        except OSError:
            _pil_cache[bold] = None
    return _pil_cache[bold]


def text_width(s, size, bold=False):
    f = _pil(bold)
    if f is None:
        return sum(size if ord(c) > 0x2E80 else size * 0.56 for c in s)
    return f.getlength(s) * size / 100.0


TAG_RE = re.compile(r'(<\/?[a-z]+(?:=[^>]*)?>)')


def strip_tags(s):
    return TAG_RE.sub('', s.replace('<br>', '\n'))


def segments(s):
    """去掉行内标记后按软换行 <br> / 换行符切分为若干行段。"""
    return strip_tags(s).split('\n')


# 中文排版禁则：行首不出现的收尾标点（允许悬挂）、行尾不出现的开头标点
_NO_START = set('，。、；：！？）」』】〕》〉”’％%,.;:!?)]}…')
_NO_END = set('（「『【〔《〈“‘([{')
_TOKEN_RE = re.compile(r'[A-Za-z0-9\.\-_/%+#&@~]+|\s|.')
FIT = 0.985   # 行宽安全系数


def layout_lines(s, width, size, bold=False, fit=FIT):
    """模拟中文排版换行，返回各行宽度：拉丁单词不拆，收尾标点可悬挂，开头标点随下一行。"""
    width *= fit
    if width <= 0:
        return [text_width(s, size, bold)]
    lines, last = [0.0], [None]
    for t in _TOKEN_RE.findall(s):
        w = text_width(t, size, bold)
        cur = lines[-1]
        if cur > 0 and cur + w > width:
            if t in _NO_START:
                lines[-1] = cur + w
                last[-1] = t
                continue
            if t.isspace():
                continue
            if last[-1] in _NO_END:
                pw = text_width(last[-1], size, bold)
                lines[-1] = cur - pw
                lines.append(pw + w)
            else:
                lines.append(w)
            last.append(t)
        else:
            lines[-1] = cur + w
            last[-1] = t
    return lines


def wrap_lines(s, width, size, bold=False):
    """估算行数（支持 <br> 软换行）。"""
    return sum(len(layout_lines(seg, width, size, bold)) for seg in segments(s))


def _orphan(lines, size, min_chars=3):
    return len(lines) >= 2 and lines[-1] < size * min_chars - 1


WARNINGS = []

# ---------------------------------------------------------------- XML 小工具


def _sub(parent, tag, **attrs):
    el = etree.SubElement(parent, qn(tag))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def _strip_style(shp):
    st = shp._element.find(qn('p:style'))
    if st is not None:
        shp._element.remove(st)
    return shp


def _set_line(shp, color=None, lw=1, dash=None, head=None, tail=None, cap=None):
    ln = shp._element.spPr.get_or_add_ln()
    for ch in list(ln):
        ln.remove(ch)
    for a in ('w', 'cap'):
        if a in ln.attrib:
            del ln.attrib[a]
    if color is None:
        _sub(ln, 'a:noFill')
        return
    ln.set('w', str(int(E(lw))))
    if cap:
        ln.set('cap', cap)
    fill = _sub(ln, 'a:solidFill')
    _sub(fill, 'a:srgbClr', val=color)
    if dash:
        _sub(ln, 'a:prstDash', val=dash)
    _sub(ln, 'a:round')
    if head:
        _sub(ln, 'a:headEnd', type=head, w='med', len='med')
    if tail:
        _sub(ln, 'a:tailEnd', type=tail, w='med', len='med')


# ---------------------------------------------------------------- 富文本
# 行内标记：<b>加粗深墨</b> <r>加粗品牌红</r> <k>仅加粗</k> <n>不加粗</n>
#          <c=RRGGBB>改色</c> <bc=RRGGBB>加粗改色</bc> <g>灰色</g> <m>等宽</m> <s=28>改字号</s>


def parse_markup(s, base):
    stack = [dict(base)]
    out = []
    for tok in TAG_RE.split(s):
        if not tok:
            continue
        m = re.fullmatch(r'<(\/?)([a-z]+)(?:=([^>]*))?>', tok)
        if m:
            closing, tag, val = m.groups()
            if tag == 'br':
                out.append(('\x0b', stack[-1]))
                continue
            if closing:
                if len(stack) > 1:
                    stack.pop()
                continue
            st = dict(stack[-1])
            if tag == 'b':
                st.update(bold=True, color=INK)
            elif tag == 'r':
                st.update(bold=True, color=RED)
            elif tag == 'k':
                st.update(bold=True)
            elif tag == 'n':
                st.update(bold=False)
            elif tag == 'c':
                st.update(color=val)
            elif tag == 'bc':
                st.update(bold=True, color=val)
            elif tag == 'g':
                st.update(color=SUB)
            elif tag == 'm':
                st.update(font=MONO)
            elif tag == 's':
                st.update(size=float(val))
            stack.append(st)
            continue
        out.append((tok, stack[-1]))
    return out


def _apply_ppr(p, align='l', lh=None, sb=0, sa=0, bullet=None, ind=0, bu_char='►', bu_size=62):
    pPr = p._p.get_or_add_pPr()
    for ch in list(pPr):
        pPr.remove(ch)
    pPr.set('algn', {'l': 'l', 'c': 'ctr', 'r': 'r', 'j': 'just', 'd': 'dist'}[align])
    if bullet or ind:
        pPr.set('marL', str(int(E(ind))))
        pPr.set('indent', str(int(E(-ind))) if bullet else '0')
    else:
        pPr.set('marL', '0')
        pPr.set('indent', '0')
    if lh:
        _sub(_sub(pPr, 'a:lnSpc'), 'a:spcPts', val=int(round(lh * 50)))
    if sb:
        _sub(_sub(pPr, 'a:spcBef'), 'a:spcPts', val=int(round(sb * 50)))
    if sa:
        _sub(_sub(pPr, 'a:spcAft'), 'a:spcPts', val=int(round(sa * 50)))
    if bullet:
        _sub(_sub(pPr, 'a:buClr'), 'a:srgbClr', val=bullet)
        _sub(pPr, 'a:buSzPct', val=int(bu_size * 1000))
        _sub(pPr, 'a:buFont', typeface=SYMBOL, pitchFamily='34', charset='0')
        _sub(pPr, 'a:buChar', char=bu_char)
    else:
        _sub(pPr, 'a:buNone')


def _add_run(p, text, st):
    if text == '\x0b':
        br = _sub(p._p, 'a:br')
        _sub(br, 'a:rPr', lang='zh-CN', altLang='en-US', sz=str(int(round(st['size'] * 50))))
        return
    r = p.add_run()
    r.text = text
    rPr = r._r.get_or_add_rPr()
    rPr.set('lang', 'zh-CN')
    rPr.set('altLang', 'en-US')
    rPr.set('sz', str(int(round(st['size'] * 50))))
    rPr.set('b', '1' if st['bold'] else '0')
    if st.get('cs'):
        rPr.set('spc', str(int(round(st['cs'] * 50))))
    fill = _sub(rPr, 'a:solidFill')
    _sub(fill, 'a:srgbClr', val=st['color'])
    _sub(rPr, 'a:latin', typeface=st['font'])
    _sub(rPr, 'a:ea', typeface=FONT)
    _sub(rPr, 'a:cs', typeface=st['font'])


def _fill_tf(tf, content, size, color, bold, font, align, lh, sb, sa, cs, bullet, ind, bu_char, bu_size):
    items = content if isinstance(content, list) else content.split('\n')
    paras = []
    for i, it in enumerate(items):
        if isinstance(it, str):
            it = {'t': it}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        st = dict(size=it.get('size', size), color=it.get('color', color), bold=it.get('bold', bold),
                  font=it.get('font', font), cs=it.get('cs', cs))
        plh = it.get('lh', lh if lh else round(st['size'] * 1.45))
        psb = it.get('sb', sb) if i > 0 else 0
        psa = it.get('sa', sa)
        pbul = it.get('bullet', bullet)
        pind = it.get('ind', ind if pbul else it.get('ind', 0))
        _apply_ppr(p, align=it.get('align', align), lh=plh, sb=psb, sa=psa, bullet=pbul, ind=pind,
                   bu_char=it.get('bu_char', bu_char), bu_size=it.get('bu_size', bu_size))
        for txt, rst in parse_markup(it['t'], st):
            _add_run(p, txt, rst)
        paras.append(dict(t=strip_tags(it['t']), segs=segments(it['t']), size=st['size'], bold=st['bold'], lh=plh,
                          sb=psb, sa=psa, ind=pind if pbul or it.get('ind') else 0, align=it.get('align', align)))
    return paras


class Deck:
    def __init__(self, total, footer_title):
        self.prs = Presentation()
        self.prs.slide_width = E(W)
        self.prs.slide_height = E(H)
        self.total = total
        self.footer_title = footer_title
        self._blank = self.prs.slide_layouts[6]
        self._patch_theme_and_layouts()
        self.slide_no = 0

    def _patch_theme_and_layouts(self):
        master = self.prs.slide_master
        theme = master.part.part_related_by(RT.THEME)
        xml = theme.blob.decode('utf-8')
        xml = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*(")', r'\g<1>微软雅黑\g<2>', xml)
        xml = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface="[^"]*"[^>]*/>\s*<a:ea typeface=")[^"]*(")',
                     r'\g<1>微软雅黑\g<2>', xml)
        theme._blob = xml.encode('utf-8')
        # 默认模板是 4:3，把母版/版式占位符横向拉伸到 16:9，方便后续手工加页
        k = W / 1440.0
        for shp in list(master.shapes) + [s for l in self.prs.slide_layouts for s in l.shapes]:
            try:
                shp.left = int(shp.left * k)
                shp.width = int(shp.width * k)
            except Exception:
                pass

    def new_slide(self):
        self.slide_no += 1
        s = self.prs.slides.add_slide(self._blank)
        s._no = self.slide_no
        return s

    def save(self, path):
        self.prs.save(path)


# ---------------------------------------------------------------- 基本图元
RECT = MSO_SHAPE.RECTANGLE
RRECT = MSO_SHAPE.ROUNDED_RECTANGLE
OVAL = MSO_SHAPE.OVAL


def shape(s, kind, x, y, w, h, fill=None, line=None, lw=1, dash=None, radius=None, rot=0, name=None):
    shp = s.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
    _strip_style(shp)
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shp.fill.background()
    _set_line(shp, line, lw, dash)
    if kind == RRECT:
        r = 6 if radius is None else radius
        shp.adjustments[0] = max(0.0, min(0.5, r / float(min(w, h))))
    if rot:
        shp.rotation = rot
    if name:
        shp.name = name
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return shp


def rect(s, x, y, w, h, fill=None, line=None, lw=1, dash=None, **kw):
    return shape(s, RECT, x, y, w, h, fill, line, lw, dash, **kw)


def rrect(s, x, y, w, h, fill=None, line=None, lw=1, dash=None, radius=6, **kw):
    return shape(s, RRECT, x, y, w, h, fill, line, lw, dash, radius=radius, **kw)


def oval(s, x, y, w, h, fill=None, line=None, lw=1, **kw):
    return shape(s, OVAL, x, y, w, h, fill, line, lw, **kw)


def _need_height(paras, avail_w, wrap=True, fit=FIT):
    need, lines_total, orphan = 0, 0, False
    for p in paras:
        n = 0
        for seg in p['segs']:
            ls = layout_lines(seg, avail_w - p['ind'], p['size'], p['bold'], fit) if wrap else [0]
            n += len(ls)
            orphan = orphan or _orphan(ls, p['size'])
        lines_total += n
        need += n * p['lh'] + p['sb'] + p['sa']
    return need, lines_total, orphan


BALANCE_FITS = (1.012, 0.975)   # 同时在“略宽”和“略窄”两种行宽假设下校验，抵消渲染器取整与字体差异


def _balance(paras, avail_w, max_frac=0.34):
    """避免末行孤字、避免“恰好放满”的临界换行：逐步收窄可用宽度，直到在各行宽假设下行数一致、
    且末行不少于 3 个字（行数不超过原先的最大值）。返回收窄量。"""
    base = [_need_height(paras, avail_w, fit=f) for f in BALANCE_FITS]
    stable = len(set(b[1] for b in base)) == 1
    if stable and not any(b[2] for b in base):
        return 0
    limit = max(b[1] for b in base)
    step = min(p['size'] for p in paras) / 2.0
    d = step
    while d <= avail_w * max_frac:
        res = [_need_height(paras, avail_w - d, fit=f) for f in BALANCE_FITS]
        if any(r[1] > limit for r in res):
            break
        if len(set(r[1] for r in res)) == 1 and not any(r[2] for r in res):
            return d
        d += step
    return 0


def text(s, x, y, w, h, content, size=22, color=BODY, bold=False, font=FONT, align='l', anchor='t', lh=None,
         sb=0, sa=0, cs=0, inset=0, wrap=True, bullet=None, ind=24, bu_char='►', bu_size=62, check=True,
         shp=None, balance=True):
    """文本框（或填充已有形状 shp 的文字）。content 为字符串（\\n 分段、<br> 软换行）或段落列表。"""
    if shp is None:
        box = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
    else:
        box = shp
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    ins = list(inset) if isinstance(inset, (tuple, list)) else [inset] * 4
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    paras = _fill_tf(tf, content, size, color, bold, font, align, lh, sb, sa, cs, bullet, ind, bu_char, bu_size)
    avail_w = w - ins[0] - ins[2]
    if wrap and balance:
        dd = _balance(paras, avail_w)
        if dd:
            if paras[0]['align'] == 'c':
                ins[0] += dd / 2.0
                ins[2] += dd / 2.0
            elif paras[0]['align'] == 'r':
                ins[0] += dd
            else:
                ins[2] += dd
            avail_w -= dd
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [E(v) for v in ins]
    if check:
        need, _, orph = _need_height(paras, avail_w, wrap)
        orph = wrap and any(_need_height(paras, avail_w, fit=f)[2] for f in BALANCE_FITS)
        avail_h = h - ins[1] - ins[3]
        if need > avail_h + 2:
            WARNINGS.append('第%02d页 文本可能溢出 (需要%.0fpx/可用%.0fpx): %s' % (
                getattr(s, '_no', 0), need, avail_h, paras[0]['t'][:30]))
        if wrap and orph:
            WARNINGS.append('第%02d页 末行孤字: %s' % (getattr(s, '_no', 0), paras[0]['t'][:30]))
        if not wrap:
            for p in paras:
                for seg in p['segs']:
                    tw = text_width(seg, p['size'], p['bold'])
                    if tw > avail_w + 2:
                        WARNINGS.append('第%02d页 单行文本超宽 (%.0f/%.0f): %s' % (getattr(s, '_no', 0), tw, avail_w,
                                                                             seg[:30]))
    return box


def line(s, x1, y1, x2, y2, color=LINE, lw=1, dash=None, head=None, tail=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    _strip_style(c)
    _set_line(c, color, lw, dash, head, tail)
    return c


def poly(s, pts, color=LINE, lw=2, dash=None, head=None, tail=None, fill=None, closed=False):
    fb = s.shapes.build_freeform(pts[0][0], pts[0][1], scale=PX)
    fb.add_line_segments(pts[1:], close=closed)
    shp = fb.convert_to_shape()
    _strip_style(shp)
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shp.fill.background()
    _set_line(shp, color, lw, dash, head, tail)
    return shp


def tri(s, cx, cy, w=26, h=34, color=TRI, direction='r'):
    """实心三角（流向）。w 为指向方向上的长度，h 为底边长。"""
    if direction in ('r', 'l'):
        shp = shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, cx - h / 2.0, cy - w / 2.0, h, w, fill=color,
                    rot=90 if direction == 'r' else 270)
    else:
        shp = shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, cx - h / 2.0, cy - w / 2.0, h, w, fill=color,
                    rot=0 if direction == 'u' else 180)
    return shp


def pic(s, path, x, y, w=None, h=None):
    return s.shapes.add_picture(path, E(x), E(y), E(w) if w else None, E(h) if h else None)


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


# ---------------------------------------------------------------- 图标（Tabler Icons, MIT）
def icon_file(name, color, stroke=1.6, filled=False):
    fn = '%s-%s-%s%s.png' % (name, color, str(stroke).replace('.', '_'), '-f' if filled else '')
    path = os.path.join(ICON_DIR, fn)
    if not os.path.exists(path):
        os.makedirs(ICON_DIR, exist_ok=True)
        spec = [dict(name=name, color=color, stroke=stroke, px=256, out=path, filled=filled)]
        tmp = os.path.join(ICON_DIR, '_spec.json')
        with open(tmp, 'w') as f:
            json.dump(spec, f)
        env = dict(os.environ)
        env.setdefault('NODE_PATH', os.path.join(TOOLS_DIR, 'node_modules'))
        subprocess.run(['node', os.path.join(TOOLS_DIR, 'render_icons.js'), tmp], check=True, env=env)
        os.remove(tmp)
    return path


def icon(s, name, color, x, y, size, stroke=1.6, filled=False):
    return pic(s, icon_file(name, color, stroke, filled), x, y, size, size)


def icon_badge(s, name, th, x, y, d=88, isz=None, fill=None, stroke=1.6):
    oval(s, x, y, d, d, fill=fill or th.soft)
    isz = isz or round(d * 0.54)
    icon(s, name, th.main, x + (d - isz) / 2.0, y + (d - isz) / 2.0, isz, stroke)


# ---------------------------------------------------------------- 页框
def frame(d, s, title, subtitle, section, title_size=54):
    rect(s, 64, 34, 1792, 2, fill=RULE)
    text(s, 64, 48, 1792, 84, title, size=title_size, color=INK, bold=True, anchor='m', lh=round(title_size * 1.3),
         wrap=False)
    rect(s, 64, 136, 96, 6, fill=RED)
    if subtitle:
        text(s, 64, 156, 1792, 42, subtitle, size=26, color=SUB, anchor='m', lh=34, wrap=False)
    footer(d, s, section)


def footer(d, s, section):
    rect(s, 64, 1010, 1792, 2, fill=RULE)
    sep = '<c=%s>   |   </c>' % RULE
    text(s, 64, 1024, 1300, 36, d.footer_title + sep + '课题申报汇报' + sep + section, size=17, color=MUTED,
         bold=True, cs=2.5, anchor='m', lh=24, wrap=False)
    text(s, 1456, 1024, 400, 36, '%02d / %02d' % (s._no, d.total), size=17, color=MUTED, bold=True, font=MONO,
         align='r', anchor='m', lh=24, wrap=False)


# ---------------------------------------------------------------- 组件
def header(s, x, y, title, note=None, w=900, size=24, color=INK, note_color=MUTED, gap=20):
    """小节标题（参考页“总体技术架构 / 三个子系统对应三个阶段”）。"""
    text(s, x, y, w, 36, title, size=size, bold=True, color=color, anchor='m', lh=round(size * 1.3), wrap=False)
    if note:
        tw = text_width(strip_tags(title), size, True)
        text(s, x + tw + gap, y, w - tw - gap, 36, note, size=19, color=note_color, bold=True, anchor='m', lh=26,
             wrap=False)


def side_label(s, x, y, w, title, desc=None, color=INK, desc_size=19):
    """左侧标签列：粗体小标题 + 灰色说明。"""
    text(s, x, y, w, 36, title, size=22, bold=True, color=color, anchor='t', lh=34)
    if desc:
        text(s, x, y + 50, w, 200, desc, size=desc_size, color=MUTED, lh=round(desc_size * 1.55))


def tag(s, x, y, txt, fill=RED, color=WHITE, size=18, h=32, padx=14, bold=True, radius=0, line=None, w=None,
        align='c'):
    """实心色块标签，返回宽度。"""
    tw = text_width(strip_tags(txt), size, bold)
    w = w or tw + padx * 2
    shp = rrect(s, x, y, w, h, fill=fill, line=line, radius=radius) if radius else rect(s, x, y, w, h, fill=fill,
                                                                                            line=line)
    text(s, x, y, w, h, txt, size=size, color=color, bold=bold, align=align, anchor='m', lh=round(size * 1.3),
         wrap=False, shp=shp, inset=(padx if align != 'c' else 0, 0, 0, 0), check=False)
    return w


def problem_card(s, x, y, w, h, title, body, title_size=27, body_size=21, pad=26, bar=GRAYBAR, bg=GRAYBG,
                 title_color=RED):
    rect(s, x, y, w, h, fill=bg, line=LINE)
    rect(s, x, y, 5, h, fill=bar)
    text(s, x + pad, y + 22, w - pad * 2, 40, title, size=title_size, bold=True, color=title_color, anchor='m',
         lh=40)
    text(s, x + pad, y + 72, w - pad * 1.6, h - 84, body, size=body_size, color=BODY, lh=round(body_size * 1.55))


def panel(s, x, y, w, h, th, title=None, num=None, icon_name=None, title_size=32, top=6, pad=40, badge=84,
          head_h=None, divider=True, bg=WHITE):
    """色条模块面板：白底 + 顶部主题色条 + 图标圆 + 标题 + 淡色编号 + 分隔线。返回内容区起点 y。"""
    rect(s, x, y, w, h, fill=bg, line=LINE)
    rect(s, x, y, w, top, fill=th.main)
    cy = y + top + 26
    head_h = head_h or badge
    tx = x + pad
    if icon_name:
        icon_badge(s, icon_name, th, x + pad, cy, badge)
        tx = x + pad + badge + 22
    if title:
        text(s, tx, cy, w - (tx - x) - pad - (110 if num else 0), head_h, title, size=title_size, bold=True,
             color=INK, anchor='m', lh=round(title_size * 1.3))
    if num:
        text(s, x + w - pad - 140, y + top + 14, 140, 70, num, size=54, bold=True, color=th.num, font=MONO,
             align='r', anchor='t', lh=70, wrap=False)
    yy = cy + head_h + 22
    if divider:
        line(s, x + pad, yy, x + w - pad, yy, LINE, 1.5)
    return yy


def bullets(s, x, y, w, h, items, th_color, size=22, lh=None, sa=10, color=BODY, ind=26, anchor='t'):
    lh = lh or round(size * 1.55)
    paras = []
    for it in items:
        if isinstance(it, str):
            it = {'t': it}
        it = dict(it)
        it.setdefault('bullet', th_color)
        paras.append(it)
    return text(s, x, y, w, h, paras, size=size, color=color, lh=lh, sa=sa, ind=ind, anchor=anchor)


def concl(s, x, y, w, h, txt, th, size=22, check=False, bar=5, pad=22, align='l', bold=True):
    """结论条：同色浅底 + 左侧色条 + 加粗同色字。"""
    rect(s, x, y, w, h, fill=th.bg)
    rect(s, x, y, bar, h, fill=th.main)
    t = ('✓ ' + txt) if check else txt
    text(s, x + bar + pad, y, w - bar - pad * 2, h, t, size=size, bold=bold, color=th.text, anchor='m',
         lh=round(size * 1.45), align=align)


def summary(s, x, y, w, h, title, desc=None, title_size=30, desc_size=21, bar=RED, desc_color=SUB, gap=10):
    """红条总结：左侧 5px 红条 + 粗体结论 + 灰色说明（内容整体垂直居中）。"""
    rect(s, x, y, 5, h, fill=bar)
    tlh = round(title_size * 1.4)
    if desc:
        dlh = round(desc_size * 1.55)
        tn = wrap_lines(title, w - 28, title_size, True)
        dn = wrap_lines(desc, w - 28, desc_size)
        total = tn * tlh + gap + dn * dlh
        top = y + max(0, (h - total) / 2.0)
        text(s, x + 28, top, w - 28, tn * tlh, title, size=title_size, bold=True, color=INK, anchor='t', lh=tlh)
        text(s, x + 28, top + tn * tlh + gap, w - 28, dn * dlh + 4, desc, size=desc_size, color=desc_color, lh=dlh)
    else:
        text(s, x + 28, y, w - 28, h, title, size=title_size, bold=True, color=INK, anchor='m', lh=tlh)


def num_mark(s, x, y, w, num, color, size=54, align='r'):
    text(s, x, y, w, size * 1.3, num, size=size, bold=True, color=color, font=MONO, align=align, anchor='t',
         lh=round(size * 1.3), wrap=False)


def ruler(s, x, y, w, color=RULER, step=44, tick=10, lw=3):
    rect(s, x, y, w, lw, fill=color)
    n = int(w // step)
    for i in range(n + 1):
        xx = x + i * step
        if xx > x + w:
            break
        rect(s, xx, y + lw, 2, tick, fill=color)


def dashed_zone(s, x, y, w, h, label=None, label_fill=ARCH, color='B7C0CE', fill=None, label_w=None, lw=1.5,
                label_size=19, label_x=None):
    rect(s, x, y, w, h, fill=fill, line=color, lw=lw, dash='dash')
    if label:
        lw_ = label_w or text_width(label, label_size, True) + 36
        lx = label_x if label_x is not None else x + (w - lw_) / 2.0
        shp = rrect(s, lx, y - 18, lw_, 36, fill=label_fill, radius=8)
        text(s, lx, y - 18, lw_, 36, label, size=label_size, bold=True, color=WHITE, align='c', anchor='m',
             lh=round(label_size * 1.3), shp=shp, wrap=False, check=False)


def layer_nav(s, x, y, w, current, bar_h=62, gap=10):
    """四层架构导航条：当前层实心高亮，其余描边弱化。返回底部 y。"""
    text(s, x, y, w, 32, '四层架构 · 当前位置', size=18, bold=True, color=MUTED, anchor='m', lh=32, cs=1)
    yy = y + 44
    for i in (4, 3, 2, 1):
        th = LAYER[i]
        on = (i == current)
        rect(s, x, yy, w, bar_h, fill=th.main if on else WHITE, line=None if on else LINE)
        if not on:
            rect(s, x, yy, 5, bar_h, fill=th.main)
        text(s, x + 20, yy + 7, w - 30, 26, LAYER_NO[i] + ' ' + LAYER_KIND[i], size=19, bold=True,
             color=WHITE if on else INK, lh=26, wrap=False)
        text(s, x + 20, yy + 33, w - 30, 24, LAYER_NAME[i], size=16, color=WHITE if on else SUB, lh=24,
             wrap=False)
        yy += bar_h + gap
    return yy


def chip(s, x, y, txt, color=BODY, fill=WHITE, line=LINE, size=19, h=38, padx=14, bold=False, radius=4, w=None,
         lw=1.2, align='c'):
    """描边小标签/节点，返回宽度。"""
    tw = text_width(strip_tags(txt), size, bold)
    w = w or tw + padx * 2
    shp = rrect(s, x, y, w, h, fill=fill, line=line, lw=lw, radius=radius) if radius else rect(
        s, x, y, w, h, fill=fill, line=line, lw=lw)
    text(s, x, y, w, h, txt, size=size, color=color, bold=bold, align=align, anchor='m', lh=round(size * 1.3),
         wrap=False, shp=shp, inset=(padx if align == 'l' else 0, 0, 0, 0), check=False)
    return w


def chips_row(s, x, y, items, gap=12, arrow=False, arrow_color=TRI, **kw):
    """横向排列 chip；arrow=True 时在 chip 之间画小三角。返回结束 x。"""
    xx = x
    for i, it in enumerate(items):
        if i:
            if arrow:
                tri(s, xx + gap / 2.0, y + kw.get('h', 38) / 2.0, 12, 16, arrow_color)
            xx += gap
        xx += chip(s, xx, y, it, **kw)
    return xx


def chips_width(items, gap=12, size=19, padx=14, bold=False):
    return sum(text_width(strip_tags(t), size, bold) + padx * 2 for t in items) + gap * (len(items) - 1)


def num_block(s, x, y, num, fill=INK, size=20, d=36, color=WHITE, font=FONT):
    shp = rect(s, x, y, d, d, fill=fill)
    text(s, x, y, d, d, num, size=size, bold=True, color=color, align='c', anchor='m', lh=round(size * 1.3),
         wrap=False, shp=shp, font=font, check=False)


def num_header(s, x, y, num, title, w=860, fill=INK, size=24, note=None):
    num_block(s, x, y, num, fill=fill)
    header(s, x + 50, y, title, note=note, w=w - 50, size=size)


def kpi(s, x, y, w, h, value, label, sub=None, color=RED, value_size=50, bg=GRAYBG, bar=None, pad=24,
        label_size=20, sub_size=18):
    rect(s, x, y, w, h, fill=bg)
    if bar:
        rect(s, x, y, 5, h, fill=bar)
    vh = round(value_size * 1.25)
    text(s, x + pad, y + 16, w - pad * 2, vh, value, size=value_size, bold=True, color=color, lh=vh, wrap=False)
    yy = y + 16 + vh + 6
    llh = round(label_size * 1.4)
    n = wrap_lines(strip_tags(label), w - pad * 2, label_size, True)
    text(s, x + pad, yy, w - pad * 2, n * llh, label, size=label_size, bold=True, color=INK, lh=llh)
    yy += n * llh + 6
    if sub:
        text(s, x + pad, yy, w - pad * 2, y + h - yy - 10, sub, size=sub_size, color=SUB, lh=round(sub_size * 1.5))


def table(s, x, y, cols, rows, header=None, size=20, lh=None, pad_x=16, pad_y=11, head_size=18, head_bg=PANEL,
          head_color=SUB, line_color=LINE, col_styles=None, row_h=None, min_row_h=0, first_line=True,
          fills=None):
    """简洁表格：表头浅底 + 行间细线。rows 中单元格可含行内标记；返回底部 y。"""
    lh = lh or round(size * 1.5)
    col_styles = col_styles or {}
    tw = sum(cols)
    yy = y
    if header:
        hh = 44
        rect(s, x, yy, tw, hh, fill=head_bg)
        cx = x
        for i, (w, h_) in enumerate(zip(cols, header)):
            st = col_styles.get(i, {})
            text(s, cx + pad_x, yy, w - 2 * pad_x, hh, h_, size=head_size, bold=True, color=head_color,
                 anchor='m', lh=round(head_size * 1.3), align=st.get('halign', st.get('align', 'l')), wrap=False)
            cx += w
        yy += hh
    elif first_line:
        line(s, x, yy, x + tw, yy, line_color, 1)
    for r_i, row in enumerate(rows):
        need = 0
        for c_i, (c, w) in enumerate(zip(row, cols)):
            st = col_styles.get(c_i, {})
            n = wrap_lines(strip_tags(c), w - 2 * pad_x, st.get('size', size), st.get('bold', False))
            need = max(need, n * st.get('lh', lh))
        rh = row_h or max(min_row_h, need + 2 * pad_y)
        if fills and fills.get(r_i):
            rect(s, x, yy, tw, rh, fill=fills[r_i])
        cx = x
        for c_i, (w, c) in enumerate(zip(cols, row)):
            st = col_styles.get(c_i, {})
            if st.get('fill'):
                rect(s, cx, yy, w, rh, fill=st['fill'])
            text(s, cx + pad_x, yy, w - 2 * pad_x, rh, c, size=st.get('size', size), color=st.get('color', BODY),
                 bold=st.get('bold', False), anchor='m', lh=st.get('lh', lh), align=st.get('align', 'l'),
                 font=st.get('font', FONT))
            cx += w
        yy += rh
        line(s, x, yy, x + tw, yy, line_color, 1)
    return yy



def _font_xml(rPr, size, color, bold):
    rPr.set('sz', str(int(round(size * 50))))
    rPr.set('b', '1' if bold else '0')
    for ch in list(rPr):
        rPr.remove(ch)
    fill = _sub(rPr, 'a:solidFill')
    _sub(fill, 'a:srgbClr', val=color)
    _sub(rPr, 'a:latin', typeface=FONT)
    _sub(rPr, 'a:ea', typeface=FONT)


def stacked_bar_chart(s, x, y, w, h, items, gap_px=2):
    """原生 100% 堆叠条形图（单条）。items: [(名称, 数值, 颜色, 条内标签或 None, 标签颜色)]。"""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    cd = CategoryChartData()
    cd.categories = ['经费']
    for name, v, *_ in items:
        cd.add_series(name, (v,))
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED_100, E(x), E(y), E(w), E(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = E(16)
    plot = ch.plots[0]
    plot.gap_width = 0
    plot.overlap = 100
    plot.vary_by_categories = False
    va, ca = ch.value_axis, ch.category_axis
    va.visible = False
    va.has_major_gridlines = False
    va.has_minor_gridlines = False
    ca.visible = False
    ca.has_major_gridlines = False
    for ser, (name, v, col, lab, lab_col) in zip(plot.series, items):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = RGBColor.from_string(col)
        ser.format.line.color.rgb = RGBColor.from_string(WHITE)
        ser.format.line.width = E(gap_px)
        if lab:
            dl = ser.points[0].data_label
            dl.has_text_frame = True
            dl.text_frame.text = lab
            dl.position = XL_LABEL_POSITION.CENTER
            for p in dl.text_frame.paragraphs:
                for r in p.runs:
                    _font_xml(r._r.get_or_add_rPr(), 18, lab_col, True)
    # 绘图区铺满图表框、图表背景透明
    cs = ch._chartSpace
    plotArea = cs.find('.//' + qn('c:plotArea'))
    layout = plotArea.find(qn('c:layout'))
    if layout is None:
        layout = etree.SubElement(plotArea, qn('c:layout'))
        plotArea.remove(layout)
        plotArea.insert(0, layout)
    for c in list(layout):
        layout.remove(c)
    ml = _sub(layout, 'c:manualLayout')
    _sub(ml, 'c:layoutTarget', val='inner')
    _sub(ml, 'c:xMode', val='edge')
    _sub(ml, 'c:yMode', val='edge')
    _sub(ml, 'c:x', val='0')
    _sub(ml, 'c:y', val='0')
    _sub(ml, 'c:w', val='1')
    _sub(ml, 'c:h', val='1')
    for parent in (cs, cs.find(qn('c:chart')).find(qn('c:plotArea'))):
        spPr = parent.find(qn('c:spPr'))
        if spPr is None:
            spPr = etree.SubElement(parent, qn('c:spPr'))
            # c:spPr 必须位于 c:txPr / c:externalData 等之前
            anchor = None
            for tagname in ('c:txPr', 'c:externalData', 'c:printSettings', 'c:userShapes', 'c:extLst'):
                anchor = parent.find(qn(tagname))
                if anchor is not None:
                    break
            if anchor is not None:
                parent.remove(spPr)
                anchor.addprevious(spPr)
        for c in list(spPr):
            spPr.remove(c)
        _sub(spPr, 'a:noFill')
        _sub(_sub(spPr, 'a:ln'), 'a:noFill')
    return gf


def stacked_bar(s, x, y, w, h, items, gap=2, size=18):
    """形状绘制的 100% 堆叠条。items: [(数值, 颜色, 条内标签或 None, 标签颜色)]。段间留 2px 白缝。"""
    total = float(sum(v for v, *_ in items))
    avail = w - gap * (len(items) - 1)
    xx = x
    for v, col, lab, fg in items:
        sw = avail * v / total
        shp = rect(s, xx, y, sw, h, fill=col)
        if lab:
            text(s, xx, y, sw, h, lab, size=size, bold=True, color=fg, align='c', anchor='m', lh=round(size * 1.3),
                 wrap=False, shp=shp)
        xx += sw + gap
    return xx
